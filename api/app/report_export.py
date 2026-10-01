"""Excel / CSV serialisation for the Reporting page. Excel mirrors RVTools'
export: one tab per report, an info tab up front (RVTools' vMetaData),
a filter on every column, frozen header row."""

import csv
import io
import zipfile
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def _excel_value(col, v):
    if v is None:
        return None
    if col.kind == "bool":
        return "Yes" if v else "No"
    if col.kind == "datetime" and isinstance(v, datetime):
        return v.astimezone(timezone.utc).replace(tzinfo=None)  # Excel has no tz; header says UTC
    return v


def _csv_value(col, v):
    if v is None:
        return ""
    if col.kind == "bool":
        return "Yes" if v else "No"
    if col.kind == "datetime" and isinstance(v, datetime):
        return v.astimezone(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")
    return v


def to_xlsx(reports: list, info_rows: list[tuple[str, str]]) -> bytes:
    """reports: [(Report, rows)]; info_rows: [(label, value)] for the Info tab."""
    wb = Workbook()
    info = wb.active
    info.title = "Info"
    for label, value in info_rows:
        info.append([label, value])
    info.column_dimensions["A"].width = 24
    info.column_dimensions["B"].width = 80
    for row in info.iter_rows(min_col=1, max_col=1):
        row[0].font = Font(bold=True)

    header_fill = PatternFill("solid", fgColor="DDE3EA")
    for report, rows in reports:
        ws = wb.create_sheet(report.label)
        cols = report.columns
        ws.append([c.label for c in cols])
        for r in rows:
            ws.append([_excel_value(c, r.get(c.key)) for c in cols])
        for i, c in enumerate(cols, 1):
            head = ws.cell(row=1, column=i)
            head.font = Font(bold=True)
            head.fill = header_fill
            head.alignment = Alignment(vertical="center")
            longest = max([len(c.label)] + [len(str(ws.cell(row=n, column=i).value or "")) for n in range(2, min(ws.max_row, 200) + 1)])
            ws.column_dimensions[get_column_letter(i)].width = min(max(longest + 2, 8), 50)
            fmt = {"num": "0.00", "int": "0", "datetime": "yyyy-mm-dd hh:mm"}.get(c.kind)
            if fmt:
                for n in range(2, ws.max_row + 1):
                    ws.cell(row=n, column=i).number_format = fmt
        ws.freeze_panes = "A2"
        if ws.max_row >= 1 and cols:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(ws.max_row, 1)}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_csv(report, rows) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([c.label for c in report.columns])
    for r in rows:
        w.writerow([_csv_value(c, r.get(c.key)) for c in report.columns])
    return buf.getvalue().encode("utf-8-sig")  # BOM so Excel opens it as UTF-8


def to_csv_zip(reports: list, stamp: str) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for report, rows in reports:
            z.writestr(f"PyXie_{report.key}_{stamp}.csv", to_csv(report, rows))
    return buf.getvalue()
